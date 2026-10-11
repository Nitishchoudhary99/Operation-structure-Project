"""End-to-end check of the C-engine workbook without Windows Excel.

usage: python tools/simulate_c_workbook.py <C-workbook> <formula-workbook> <harness> <work-dir> <recalc.py>

1. exports the C workbook's input blocks (full capacity, e.g. 20,000 rows)
2. runs the C engine through the test harness (exactly what Excel would call)
3. pastes the engine output over the six array-formula blocks, as Excel would display them
4. recalculates every downstream formula (Dashboard, Reports, Invoice) in LibreOffice
5. compares Dashboard / Reports / Invoice with the formula edition, cell by cell
"""
import json
import math
import os
import subprocess
import sys

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

c_path, f_path, harness, work, recalc = sys.argv[1:6]
os.makedirs(work, exist_ok=True)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from verify_engine_io import enc, read_out, serial  # noqa: E402

wb = load_workbook(c_path)
blocks = {"items": ("IN_Items",), "parties": ("IN_Parties",), "purchases": ("IN_Purchases",),
          "sales": ("IN_Sales",), "cashbook": ("IN_CashBook",), "adjust": ("IN_Adjust",), "settings": ("IN_Settings",)}
vals = load_workbook(c_path, data_only=True)


def rng(defname):
    sheet, ref = wb.defined_names[defname].attr_text.split("!")
    a, b = ref.replace("$", "").split(":")
    return sheet.strip("'"), a, b


def split(ref):
    col = "".join(ch for ch in ref if ch.isalpha())
    return col, int(ref[len(col):])


for key, (defname,) in blocks.items():
    sheet, a, b = rng(defname)
    (c1, r1), (c2, r2) = split(a), split(b)
    ws = vals[sheet]
    rows = [[ws.cell(r, c).value for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1)]
            for r in range(r1, r2 + 1)]
    if key == "settings":   # B10 is =TODAY() in the file; pin it to the formula edition's as-of date
        rows[6][0] = load_workbook(f_path, data_only=True)["Settings"]["B10"].value
    with open(os.path.join(work, f"in_{key}.tsv"), "w", encoding="utf-8") as f:
        f.write(f"{len(rows)} {len(rows[0])}\n")
        for row in rows:
            f.write("\t".join(enc(v) for v in row) + "\n")

out = subprocess.run([harness, "verify", work], capture_output=True, text=True)
print(out.stdout.strip())
if out.returncode:
    sys.exit(out.stderr)

# paste engine output where Excel would show it
targets = [("Items", "L", "items"), ("Purchases", "I", "purchases"), ("Sales", "J", "sales"),
           ("Parties", "J", "parties"), ("CashBook", "I", "cashbook"), ("Adjustments", "F", "adjust")]
for sheet, c1, key in targets:
    ws = wb[sheet]
    data = read_out(work, key)
    base = column_index_from_string(c1)
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            ws.cell(5 + r, base + c).value = None if v == "" else v
wb["Settings"]["B10"] = load_workbook(f_path, data_only=True)["Settings"]["B10"].value
wb["Dashboard"]["K4"] = "simulated"
sim = os.path.join(work, "simulated.xlsx")
wb.save(sim)
res = json.loads(subprocess.run([sys.executable, recalc, sim, "600"], capture_output=True, text=True).stdout)
print(f"LibreOffice recalc of simulated workbook: {res.get('status')}, {res.get('total_formulas')} formulas, "
      f"{res.get('total_errors')} errors")

got = load_workbook(sim, data_only=True)
exp = load_workbook(f_path, data_only=True)
checks = [("Dashboard", "B8:N15"), ("Dashboard", "B20:G24"), ("Dashboard", "I20:N29"),
          ("Reports", "A5:M17"), ("Reports", "A23:F30"), ("Reports", "H23:J34"), ("Invoice", "A7:K42")]
total = bad = 0
for sheet, ref in checks:
    for erow, grow in zip(exp[sheet][ref], got[sheet][ref]):
        for e, g in zip(erow, grow):
            total += 1
            ev, gv = serial(e.value), serial(g.value)
            ok = (ev == gv) or (isinstance(ev, (int, float)) and isinstance(gv, (int, float))
                                and math.isclose(ev, gv, rel_tol=1e-9, abs_tol=0.0051))
            if not ok and ev in (None, "") and gv in (None, ""):
                ok = True
            # Stock Value card: the C engine values each item to the paisa before totalling
            # (formula edition sums unrounded fractions) -> allow 0.5 paisa per item
            if not ok and sheet == "Dashboard" and e.coordinate == "H12":
                n_items = sum(1 for row in exp["Items"]["A5:A504"] if row[0].value)
                ok = abs(ev - gv) <= 0.005 * n_items
                print(f"  Dashboard!H12 stock value: formula {ev:.4f} vs C {gv:.2f} "
                      f"(paisa rounding per item, allowed ±{0.005 * n_items:.2f}) -> {'OK' if ok else 'FAIL'}")
            if not ok:
                bad += 1
                if bad <= 10:
                    print(f"  MISMATCH {sheet}!{e.coordinate}: formula edition={e.value!r} C edition={g.value!r}")
print(f"Dashboard / Reports / Invoice: {total} cells compared, mismatches: {bad}")
sys.exit(1 if bad or res.get("total_errors") else 0)
